"""Teacher materials of each topic and skill, under teacher/unit-N/<tid>/[<aid>/].

<skill>/problem/solution.py   -> solution.json, solution.html, <aid>-solution.ipynb, solution-preview.html
<skill>/problem/              -> <aid>-student.ipynb, student-preview.html, assign.html
<skill>/atom.md + tutor.md    -> pre-lecture-post.html, post-lecture-post.html, <aid>-{pre,post}-lecture-handout.docx
<skill>/activity.md           -> teacher/unit-N/<tid>/activity-keys.html
all posts                     -> teacher/feed.xml
"""
import hashlib
import html
import io
import json
import re
from string import Template

from . import course as C
from . import nav as N
from .atomtext import bullets, md, one_line, parse_atom, sections, web_code, web_prompt, web_tests, web_text
from .course import BRAND, SITE_URL
from .paths import ASSETS, TEACHER, TEMPLATES, teacher_atom, teacher_topic
from .util import dump_json, emit, emit_bytes, esc, rd, rel

TUTOR_CSS = ASSETS / "teacher" / "tutor-theme.css"

PAGE_HEAD = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title}</title>
<link rel="stylesheet" href="{css}">
</head>
<body>
"""


# ------------------------------------------------------------- solutions --
def solutions():
    for a in C.atom_order():
        p = a["src"] / "problem"
        if not (p / "solution.py").exists():
            continue
        d = teacher_atom(a["id"])
        emit(d / "solution.json", dump_json({"atom_id": a["id"], "topic_id": C.tid_of(a["id"]), "title": a["title"],
                                             "solution_py": web_code(rd(p / "solution.py"))}), json_equal=True)
        path = d / "solution.html"
        text = Template(rd(TEMPLATES / "pages" / "solution.html")).substitute(
            assets=rel(d, ASSETS), aid=a["id"], tid=C.tid_of(a["id"]), title=esc(a["title"]),
            topic_page=rel(d, teacher_topic(C.tid_of(a["id"])) / "index.html"), problem_page=rel(d, N.problem_page(a)))
        emit(path, text)


# ----------------------------------------------------------- tutor posts --
POST_TAIL = """</pre>
</div>
<script>
function toggleMode() {
  var rendered = document.getElementById("rendered");
  var source = document.getElementById("source");
  var btn = document.getElementById("mode-toggle");
  var showingSource = source.style.display !== "none";
  rendered.style.display = showingSource ? "block" : "none";
  source.style.display = showingSource ? "none" : "block";
  btn.textContent = showingSource ? "View Source (Markdown)" : "View Rendered";
}
</script>
</body>
</html>
"""


def post_markdown(aid, a, t, kind):
    if kind == "pre":
        parts = [f"![banner](banner.png)\n\n## {aid}, {a['title']}\n**Before class today.** {a['lead']}",
                 f"**The intuition to hold onto:** {t['The intuition to hold onto']}"]
        if a["misconceptions"]:
            parts.append(f"**A myth worth busting before we start:** {a['misconceptions'][0]}")
        parts.append(f"**Today's goal:** {a['objective']}\nHandout: `pre-lecture-handout.docx`.")
    else:
        parts = [f"![banner](banner.png)\n\n## {aid} recap, {a['title']}\nNice work today. Here's what to lock in:",
                 "- Recap handout: `post-lecture-handout.docx`\n- Simulation: SimLab\n- Self-check quiz: due before next class",
                 f"**The big picture:** {t['The big picture']}"]
        if a["terms"]:
            parts.append("**Key terms to lock in:**\n" + "\n".join(f"- {x}" for x in a["terms"]))
        if a["misconceptions"]:
            parts.append("**Common misconceptions -- don't fall for these:**\n" + "\n".join(f"- {x}" for x in a["misconceptions"]))
        parts.append(f"**Learning objective(s):** {a['los']}")
    return "\n\n".join(parts) + "\n"


def posts():
    feed = []
    for atom in C.atom_order():
        aid = atom["id"]
        a = parse_atom(atom["src"])
        t = dict(sections(rd(atom["src"] / "tutor.md"))[1])
        d = teacher_atom(aid)
        for kind, label in (("pre", "Pre-Lecture"), ("post", "Post-Lecture")):
            src = post_markdown(aid, a, t, kind)
            page = (PAGE_HEAD.format(title=f"{aid} {label} Post · {BRAND}", css=rel(d, TUTOR_CSS))
                    + f'<nav class="top-nav">\n  <div class="crumb"><a href="{rel(d, teacher_topic(C.tid_of(aid)) / "index.html")}">'
                      f'{C.tid_of(aid)} teacher page</a> &gt; {aid} &gt; {label}</div>\n'
                      '  <button class="mode-btn" id="mode-toggle" onclick="toggleMode()">View Source (Markdown)</button>\n</nav>\n'
                      f'<div class="post-wrap">\n  <div class="post-rendered" id="rendered">\n{md(src)}\n  </div>\n'
                      f'  <pre class="post-source" id="source" style="display:none;">{html.escape(src, quote=False)}'
                    + POST_TAIL)
            path = d / f"{kind}-lecture-post.html"
            emit(path, page)
            feed.append((f"{label}: {aid}, {a['title']}", path.relative_to(TEACHER).as_posix()))
    base = SITE_URL + "teacher/"
    items = "".join(f"    <item><title>{esc(t)}</title><link>{base}{p}</link><guid>{base}{p}</guid></item>\n" for t, p in feed)
    emit(TEACHER / "feed.xml", '<?xml version="1.0" encoding="UTF-8"?>\n<rss version="2.0"><channel>\n'
         f"  <title>{BRAND} Cybersecurity, Tutor Posts</title>\n  <link>{SITE_URL}</link>\n"
         "  <description>Pre/post-lecture classroom posts, in course order.</description>\n"
         f"{items}</channel></rss>\n")


# ------------------------------------------------------------- notebooks --
NB_RUNNER = '''

# ---- run every test above ------------------------------------------------
_results = []
for _name, _fn in list(globals().items()):
    if _name.startswith("test_") and callable(_fn):
        try:
            _fn()
            _results.append((_name, None))
        except Exception as _e:
            _results.append((_name, _e))
for _name, _e in _results:
    print(("PASS  " if _e is None else "FAIL  ") + _name + ("" if _e is None else f": {_e!r}"))
print(f"{sum(e is None for _, e in _results)}/{len(_results)} tests passed")
'''

NB_STYLE = """<style>
.nb-cell { margin-bottom: 14px; }
.nb-md { font-family: var(--font-serif); font-size: 15.5px; }
.nb-code { font-family: var(--font-mono); font-size: 13px; background: #1e1e1e; color: #d4d4d4; padding: 14px 16px; white-space: pre-wrap; overflow-x: auto; }
body { padding: 20px 26px; }
</style>
"""

ASSIGN_STYLE = """<style>
.assign-wrap { max-width: 900px; margin: 0 auto; padding: 30px; }
.assign-actions { display: flex; gap: 12px; margin: 18px 0 24px; flex-wrap: wrap; }
.assign-btn { font-family: var(--font-sans); font-size: 13px; text-transform: uppercase; letter-spacing: .5px; padding: 10px 18px; text-decoration: none; font-weight: 600; }
.assign-btn.primary { background: var(--accent-color); color: #fff; }
.assign-btn.primary:hover { background: var(--accent-dark); }
.assign-btn.secondary { background: #fff; color: var(--accent-color); border: 1px solid var(--accent-color); }
.assign-frame { width: 100%; height: 640px; border: 1px solid var(--border-color); background: #fff; }
</style>
"""


def notebook_cells(a, solution):
    p = a["src"] / "problem"
    heading = f"# {a['title']}" + ("  (MODEL SOLUTION, instructor copy, do not distribute)" if solution else "")
    intro = heading + "\n\n" + web_prompt(rd(p / "prompt.md")) + "\n## Reference\n" + \
        re.sub(r"(?m)^#( |#)", r"##\1", web_text(rd(p / "interface.md"))).split("\n", 1)[-1].lstrip("\n")
    code = ("# Reference solution\n" + web_code(rd(p / "solution.py"))) if solution else \
        ("# Your implementation\n" + web_code(rd(p / "starter.py")))
    tests = web_tests(rd(p / "tests.py")).replace("import impl\n", "import __main__ as impl  # the functions defined above\n", 1)
    tests = tests.rstrip("\n") + "\n" + NB_RUNNER
    return [("markdown", intro), ("code", code),
            ("markdown", "## Self-check tests\nRun the cell below to check your `solve` implementation above."),
            ("code", tests)]


def ipynb(aid, cells, salt):
    nb = {"cells": [], "metadata": {"kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
                                    "language_info": {"name": "python"}},
          "nbformat": 4, "nbformat_minor": 5}
    for i, (kind, src) in enumerate(cells):
        cid = hashlib.sha1(f"{aid}-{salt}-{i}".encode()).hexdigest()[:8]
        c = {"cell_type": kind, "id": cid, "metadata": {}, "source": src.splitlines(keepends=True)}
        if kind == "code":
            c = {"cell_type": kind, "execution_count": None, "id": cid, "metadata": {}, "outputs": [],
                 "source": src.splitlines(keepends=True)}
        nb["cells"].append(c)
    return json.dumps(nb, ensure_ascii=False, indent=1) + "\n"


def preview(cells, title, css):
    body = "".join(
        f'<div class="nb-cell nb-md">{md(src)}</div>\n' if kind == "markdown"
        else f'<div class="nb-cell nb-code">{html.escape(src, quote=False)}</div>\n'
        for kind, src in cells)
    return PAGE_HEAD.format(title=title, css=css).replace("</head>", NB_STYLE + "</head>") + body + "</body>\n</html>\n"


def notebooks():
    for a in C.atom_order():
        if not N.has_problem(a):
            continue
        aid, d = a["id"], teacher_atom(a["id"])
        css = rel(d, TUTOR_CSS)
        cells = notebook_cells(a, solution=False)
        emit(d / f"{aid}-student.ipynb", ipynb(aid, cells, "student"))
        emit(d / "student-preview.html", preview(cells, f"{a['title']}, Student", css))
        emit(d / "assign.html",
             PAGE_HEAD.format(title=f"Assignment {aid} · {BRAND}", css=css).replace("</head>", ASSIGN_STYLE + "</head>")
             + f'<nav class="top-nav">\n  <div class="crumb"><a href="{rel(d, teacher_topic(C.tid_of(aid)) / "index.html")}">'
               f'{C.tid_of(aid)} teacher page</a> &gt; Assignment &gt; {aid}</div>\n</nav>\n'
               f'<div class="assign-wrap">\n  <h1 class="hub-title">{a["title"]}</h1>\n  <p class="hub-subtitle">Skill {aid} &middot; Coding Assignment</p>\n'
               '  <div class="assign-actions">\n'
               f'    <a class="assign-btn primary" href="{rel(d, N.problem_page(a))}">Open Live Playground &rarr;</a>\n'
               f'    <a class="assign-btn secondary" href="{aid}-student.ipynb" download>Download Notebook (.ipynb)</a>\n  </div>\n'
               '  <iframe class="assign-frame" src="student-preview.html"></iframe>\n</div>\n</body>\n</html>\n')
        scells = notebook_cells(a, solution=True)
        emit(d / f"{aid}-solution.ipynb", ipynb(aid, scells, "solution"))
        emit(d / "solution-preview.html", preview(scells, f"{a['title']}, Model Solution", css))


# -------------------------------------------------------------- handouts --
INK = (0x1B, 0x1B, 0x3A)


def docx_runs(par, text):
    """Add markdown inline text (**bold**, *italic*, `code`) to a python-docx paragraph."""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)", one_line(text)):
        if not tok:
            continue
        if tok.startswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Consolas"
        elif tok.startswith("*"):
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok)


def handout(aid, a, kind):
    import docx
    from docx.shared import RGBColor
    doc = docx.Document()
    for st in ("Title", "Heading 2"):
        doc.styles[st].font.color.rgb = RGBColor(*INK)
    doc.core_properties.title = f"{aid} {kind}-lecture handout"
    doc.add_paragraph(f"{aid}, {a['title']}", style="Title")

    def para(text):
        docx_runs(doc.add_paragraph(), text)

    def items(xs):
        for x in xs:
            docx_runs(doc.add_paragraph(style="List Bullet"), x)

    if kind == "pre":
        para(a["unit_line"])
        para(f"AP Learning Objective(s): {a['los']}   |   Skill: {a['skill']}")
        doc.add_heading("Before class", level=2)
        para(a["lead"])
        doc.add_heading("Terms you'll see today", level=2)
        items([re.match(r"\*\*(.+?)\*\*", x).group(1) if x.startswith("**") else x.split(",")[0] for x in a["terms"]])
        doc.add_heading("Come ready to", level=2)
        para("Think about a time a message tried to rush, flatter, scare, or pressure you into acting quickly. "
             "Be ready to describe it in one sentence." if aid.startswith("1.1.") else
             "Skim the terms above and write down one question you want answered in class.")
    else:
        para(f"{a['unit_line']}  ·  Recap / take-home reference")
        doc.add_heading("Lesson recap", level=2)
        for p in a["lesson_paras"]:
            if p.startswith("### "):
                doc.add_heading(p[4:].split("\n")[0], level=3)
                p = p.split("\n", 1)[1] if "\n" in p else ""
            if re.match(r"^(- |\d+\. )", p):
                items([re.sub(r"^\d+\. ", "", x) for x in bullets(re.sub(r"(?m)^\d+\. ", "- ", p))])
            elif p.startswith("```"):
                r = doc.add_paragraph().add_run(p.strip("`").strip("\n"))
                r.font.name = "Consolas"
            elif p:
                para(p)
        doc.add_heading("Key terms", level=2)
        items(a["terms"])
        doc.add_heading("Common misconceptions", level=2)
        items(a["misconceptions"])
        doc.add_heading("Try it", level=2)
        items(a["try"])
        doc.add_heading("Self-check", level=2)
        para("Complete the quiz for this skill before the next class.")
    return doc


def docx_signature(blob):
    """Text and style of every paragraph: Word files differ byte-wise on every save, so compare these."""
    import docx
    d = docx.Document(io.BytesIO(blob))
    return [(p.style.name, p.text, tuple((r.bold, r.italic) for r in p.runs)) for p in d.paragraphs]


def handouts():
    for atom in C.atom_order():
        a = parse_atom(atom["src"])
        for kind in ("pre", "post"):
            buf = io.BytesIO()
            handout(atom["id"], a, kind).save(buf)
            emit_bytes(teacher_atom(atom["id"]) / f"{atom['id']}-{kind}-lecture-handout.docx", buf.getvalue(),
                       same=lambda old, new: docx_signature(old) == docx_signature(new))


# ---------------------------------------------------------- activity keys --
def activity_keys():
    from .materials import activity
    for tid in C.order():
        keys = [(a, key) for a in C.topic_atoms(tid) if (a["src"] / "activity.md").exists() for key in [activity(a)[1]] if key]
        if not any((a["src"] / "activity.md").exists() for a in C.topic_atoms(tid)):
            continue
        d = teacher_topic(tid)
        k = len(keys)
        body = "".join(f'<div class="poster">\n  <div class="poster-meta">Skill {a["id"]}</div>\n  <h1>{a["title"]}</h1>\n'
                       + "".join(f"  <h2>{t}</h2>\n{md(b)}\n" for t, b in key) + "</div>\n" for a, key in keys)
        emit(d / "activity-keys.html",
             PAGE_HEAD.format(title=f"{tid} Activity Answer Keys · {BRAND}", css=rel(d, ASSETS / "activities" / "activity-theme.css"))
             + '<nav class="top-nav">\n  <div class="crumb"><a href="index.html">' + f'{tid} teacher page</a> &gt; '
               'Activity answer keys</div>\n  <button class="print-btn" onclick="window.print()">Print All</button>\n</nav>\n'
               f'<div class="hub-header">\n  <h1 class="hub-title">Topic {tid} Activity Answer Keys</h1>\n'
               f'  <p class="hub-subtitle">{C.topic(tid)["title"]} &middot; instructor-only &middot; {k} entr{"y" if k == 1 else "ies"}</p>\n</div>\n'
             + body + "</body>\n</html>\n")


def run():
    solutions()
    posts()
    notebooks()
    handouts()
    activity_keys()
