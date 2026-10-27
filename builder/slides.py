"""Topic slide decks.

source/units/.../<topic>/slides.json   -> unit-N/<tid>/slides.html          (student deck)
                                       -> teacher/unit-N/<tid>/slides.html  (with speaker notes)
                                       -> downloads/unit-N/Topic_<tid>_Slides.pptx (classroom PowerPoint)
source/units/.../<topic>/images/<diagram_ref>.png   the slides' diagrams (Mermaid code beside them as .mmd)
templates/topic_slides.html is the page shell.

A topic file holds the deck (slides with beats 1-3: bullets, explanation paragraphs, diagram) and, under
"classroom", the text used only by the PowerPoint: cover summary, learning objectives, concept maps (stored
as their PowerPoint shape XML) and the Big Idea slide.
"""
import html
import io
import json
import re

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from . import course as C
from . import nav as N
from .paths import ASSETS, TEMPLATES, downloads_unit, teacher_topic, topic_dir
from .util import emit, emit_bytes, rel

TEMPLATE = TEMPLATES / "topic_slides.html"


def diagram_png(tid, ref):
    return C.topic(tid)["src"] / "images" / f"{ref}.png" if ref else None


# ------------------------------------------------------------ HTML decks --
def deck_data(d, tid, audience, page):
    back = rel(page.parent, N.target(tid, "topic"))
    out = {k: d[k] for k in ("subunit_id", "subunit_title", "unit", "unit_title", "learning_objectives")}
    out["audience"] = audience
    out["landing_href"] = back
    out["slides"] = []
    for s in d["slides"]:
        png = diagram_png(tid, s.get("diagram_ref"))
        src = rel(page.parent, topic_dir(tid) / "images" / png.name) if png and png.exists() else None
        rec = {}
        for k, v in s.items():                        # source order, image_src right after diagram_prompt
            if k == "speaker_notes" and audience == "student":
                continue
            rec[k] = v
            if k == "diagram_prompt":
                rec["image_src"] = src
        rec.setdefault("image_src", src)
        out["slides"].append(rec)
    return out, back


def deck_html(d, tid, audience, page):
    data, back = deck_data(d, tid, audience, page)
    blob = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    return (TEMPLATE.read_text(encoding="utf-8")
            .replace("__TITLE__", html.escape(f'{d["subunit_id"]}, {d["subunit_title"]}'))
            .replace("__ASSETS__", rel(page.parent, ASSETS))
            .replace("__BACK__", back).replace("__TID__", tid).replace("__DATA__", blob, 1))


# ------------------------------------------------------- classroom pptx --
ACCENT, DEEP, DARK = RGBColor(0xD7, 0x26, 0x3D), RGBColor(0x9C, 0x1C, 0x2C), RGBColor(0x6E, 0x14, 0x1F)
INK, BODY, GREY = RGBColor(0x1A, 0x1A, 0x1A), RGBColor(0x3F, 0x3F, 0x3F), RGBColor(0x7A, 0x7A, 0x7A)
SOFT, LINE, FAINT, WHITE = RGBColor(0xFB, 0xEE, 0xF0), RGBColor(0xE9, 0xE2, 0xE3), RGBColor(0xFC, 0xE4, 0xE7), RGBColor(0xFF, 0xFF, 0xFF)
SANS, SERIF = "Calibri", "Cambria"
INLINE = re.compile(r"(<b>.*?</b>|<strong>.*?</strong>|<i>.*?</i>|<em>.*?</em>|<code>.*?</code>)", re.S)


def plain(h):
    return " ".join(html.unescape(re.sub(r"<[^>]+>", " ", h or "")).split())


def box(slide, x, y, w, h, text="", size=12, color=BODY, bold=False, italic=False, font=SANS,
        align=None, anchor=MSO_ANCHOR.TOP, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    if fill is not None:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    if line is not None:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(45720)
    tf.margin_top = tf.margin_bottom = Emu(22860)
    if text:
        runs(tf.paragraphs[0], text, size, color, bold, italic, font)
        tf.paragraphs[0].alignment = align if align is not None else PP_ALIGN.LEFT
    return shp


def runs(par, h, size, color, bold=False, italic=False, font=SANS):
    """Inline HTML (<b>, <i>, <code>) or plain text -> runs."""
    for tok in INLINE.split(h if "<" in (h or "") else html.escape(h or "", quote=False)):
        if not tok:
            continue
        m = re.match(r"<(b|strong|i|em|code)>(.*)</\1>$", tok, re.S)
        b, i, f, inner = bold, italic, font, tok
        if m:
            tag, inner = m.groups()
            b, i = b or tag in ("b", "strong"), i or tag in ("i", "em")
            f = "Consolas" if tag == "code" else font
        t = html.unescape(re.sub(r"<[^>]+>", "", inner))
        if not t:
            continue
        r = par.add_run()
        r.text = t
        r.font.name, r.font.size, r.font.bold, r.font.italic = f, Pt(size), b, i
        r.font.color.rgb = color


def notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def header(slide, kicker, title, quiz=False, title_size=26):
    box(slide, 457200, 365760, 7315200, 274320, kicker, 10, GREY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if quiz:
        box(slide, 11441887, 384048, 292608, 292608, "?", 12, WHITE, bold=True, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE, fill=ACCENT, shape=MSO_SHAPE.OVAL)
    else:
        box(slide, 11569903, 384048, 164592, 164592, fill=ACCENT, shape=MSO_SHAPE.OVAL)
    box(slide, 457200, 658368, 11277295, 777240, title, title_size, INK, bold=True, font=SERIF, anchor=MSO_ANCHOR.MIDDLE)


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    return s


def cover(prs, d, c):
    s = blank(prs)
    box(s, 548640, 1600200, 10058400, 320040, "AAYAT-AP-PREP  ·  CYBERSECURITY", 12, ACCENT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 548640, 1965960, 10515600, 365760, d["unit_title"].upper(), 14, GREY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 548640, 2423160, 11064240, 1554480, f'Topic {d["subunit_id"]}: {d["subunit_title"]}', 40, INK, bold=True, font=SERIF)
    box(s, 548640, 4160520, 9875520, 1737360, c["summary"], 14, BODY)
    box(s, 548640, 6263640, 11064240, 320040, c["footer"], 10.5, GREY, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, c["notes"])


def objectives(prs, d, o):
    s = blank(prs)
    box(s, 457200, 411480, 7315200, 274320, f'TOPIC {d["subunit_id"]}', 10, GREY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 457200, 685800, 9144000, 640080, "Learning Objectives", 30, INK, bold=True, font=SERIF, anchor=MSO_ANCHOR.MIDDLE)
    items = o["items"]
    top, bottom = 1600200, 6400800
    gap = 200000
    h = min(1466088, (bottom - top - gap * (len(items) - 1)) // max(len(items), 1))
    for k, it in enumerate(items):
        y = top + k * (h + gap)
        box(s, 457200, y, 11277295, h, fill=SOFT, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        box(s, 640080, y + (h - 512064) // 2, 1051560, 512064, it["code"], 12, WHITE, bold=True, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE, fill=DEEP, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        t = box(s, 1874520, y, 9722815, h, it["skill"], 11, DEEP, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        par = t.text_frame.add_paragraph()
        par.alignment = PP_ALIGN.LEFT
        runs(par, it["text"], 13, BODY)
    if o.get("footer"):
        box(s, 457200, 6492240, 11277295, 274320, o["footer"], 9, GREY, italic=True, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, o["notes"])


def concept_map(prs, d, m):
    s = blank(prs)
    header(s, f'TOPIC {d["subunit_id"]}   ·   CONCEPT MAP', m["title"])
    box(s, 457200, 1440000, 800000, 34290, fill=ACCENT)
    tree = s.shapes._spTree
    next_id = max(int(x.get("id")) for x in tree.iter() if x.tag.endswith("}cNvPr")) + 1
    for xml in m["shapes_xml"]:
        el = etree.fromstring(xml)
        for c in el.iter():
            if c.tag.endswith("}cNvPr"):
                c.set("id", str(next_id))
                next_id += 1
        tree.append(el)
    notes(s, m["notes"])


def big_idea(prs, d, b):
    s = blank(prs)
    box(s, 8229600, -1463040, 5852160, 5852160, fill=SOFT, shape=MSO_SHAPE.OVAL)
    box(s, 640080, 594360, 6400800, 274320, f'TOPIC {d["subunit_id"]}', 10, GREY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 640080, 914400, 7315200, 640080, "The Big Idea", 30, INK, bold=True, font=SERIF, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 640080, 1920240, 6492240, 3291840, b["text"], 22, BODY, italic=True, font=SERIF, anchor=MSO_ANCHOR.MIDDLE)
    if b.get("closing"):
        box(s, 640080, 5989320, 7863840, 457200, b["closing"], 12, GREY, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, b["notes"])


def divider(prs, d, n, title):
    s = blank(prs)
    box(s, 0, 822960, 12191695, 3931920, f"{n:02d}", 220, FAINT, bold=True, font=SERIF, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE)
    box(s, 1371600, 3063240, 9445752, 914400, title, 27, INK, bold=True, font=SERIF, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE)
    box(s, 548640, 6309360, 11064240, 320040, f'Topic {d["subunit_id"]}', 10.5, GREY, bold=True, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE)
    notes(s, f"Transition to part {n} of the topic. Pause here for questions before continuing.")


def content(prs, d, tid, group):
    first = group[0]
    beats = {s["beat"]: s for s in group}
    s = blank(prs)
    header(s, f'TOPIC {d["subunit_id"]}   ·   {first["index"]}', plain(first["title"]), quiz=bool(first.get("quiz")))
    bullets = (beats.get(1) or first).get("bullets") or []
    if not bullets:                                   # slides with paragraphs only: show the first paragraphs
        bullets = [b for b in (beats.get(1) or first).get("content_blocks", []) if isinstance(b, str)][:3]
    ref = next((x.get("diagram_ref") for x in group if x.get("diagram_ref")), "")
    png = diagram_png(tid, ref)
    has_img = bool(png and png.exists())
    top, bottom = 1554480, 6400800
    width = 4407280 if has_img else 10900000
    n = max(len(bullets), 1)
    step = (bottom - top) // n
    size = 13 if n <= 5 else 11
    for k, bl in enumerate(bullets):
        y = top + k * step
        box(s, 457200, y + (step - 201168) // 2 if step < 900000 else y + 310896, 201168, 201168, fill=DEEP, shape=MSO_SHAPE.OVAL)
        box(s, 786384, y, width, step, bl, size, BODY, anchor=MSO_ANCHOR.MIDDLE)
    if has_img:
        from PIL import Image
        with Image.open(png) as im:
            iw, ih = im.size
        rx, ry, rw, rh = 5513704, 1554480, 6220791, 4562856
        sc = min(rw / iw, rh / ih)
        pw, ph = int(iw * sc), int(ih * sc)
        s.shapes.add_picture(str(png), Emu(rx + (rw - pw) // 2), Emu(ry + (rh - ph) // 2), Emu(pw), Emu(ph))
        box(s, rx, 6163056, rw, 256032, ref.replace("_", " ").title(), 9, GREY, bold=True, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
    explain = beats.get(2)
    if explain:
        notes(s, "\n\n".join(plain(b) for b in explain.get("content_blocks", []) if isinstance(b, str)))
    elif first.get("speaker_notes"):
        notes(s, plain(first["speaker_notes"]))


def classroom_pptx(d, tid):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    c = d["classroom"]
    cover(prs, d, c["cover"])
    objectives(prs, d, c["objectives"])
    for m in c["concept_maps"]:
        concept_map(prs, d, m)
    big_idea(prs, d, c["big_idea"])
    groups, sections = {}, {}
    for sl in d["slides"]:
        groups.setdefault(sl["index"], []).append(sl)
    last_section = None
    for idx, group in groups.items():
        section = idx.split(".")[2] if idx.count(".") >= 2 else idx
        if section != last_section:
            sections[section] = len(sections) + 1
            divider(prs, d, sections[section], plain(group[0]["title"]))
            last_section = section
        content(prs, d, tid, group)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), len(prs.slides)


def pptx_signature(blob):
    """Every part of the package except the timestamped core properties, so an unchanged deck is not
    rewritten just because the zip was re-created."""
    import zipfile
    z = zipfile.ZipFile(io.BytesIO(blob))
    return {n: z.read(n) for n in z.namelist() if n != "docProps/core.xml"}


def run(only=()):
    for tid in C.order():
        src = C.topic(tid)["src"] / "slides.json"
        if not src.exists() or (only and tid not in only):
            continue
        d = json.loads(src.read_text(encoding="utf-8"))
        for audience, page in (("student", N.target(tid, "slides")), ("instructor", teacher_topic(tid) / "slides.html")):
            emit(page, N.apply(page, deck_html(d, tid, audience, page)))
        blob, n = classroom_pptx(d, tid)
        emit_bytes(downloads_unit(tid) / f"Topic_{tid}_Slides.pptx", blob, same=lambda a, b: pptx_signature(a) == pptx_signature(b))
