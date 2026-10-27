"""
Write a lecture deck as PowerPoint (used by lectures.py).

One PowerPoint slide per browser-deck slide, so slide numbers match across the two formats.
Blocks go where the browser deck draws them (positions measured in Chromium, scaled to the
slide); without a browser they are stacked by estimated height. Theme and layouts come from
templates/lecture.pptx. A list/steps/flow block with `smartart=<Layout>` becomes real SmartArt
built from templates/smartart/<Layout>/ (layout, style and colour definitions), with the
block's items as its text.
"""
import copy
import html
import json
import re
import uuid
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .paths import TEMPLATES
SMARTART = TEMPLATES / "smartart"

RED, RED_DEEP, RED_SOFT = RGBColor(0xB3, 0x00, 0x1B), RGBColor(0x8A, 0x00, 0x15), RGBColor(0xFB, 0xEC, 0xEE)
INK, GREY, LINE, CODE_BG = RGBColor(0x20, 0x1C, 0x1A), RGBColor(0x6E, 0x66, 0x60), RGBColor(0xE7, 0xE1, 0xDA), RGBColor(0x1E, 0x1E, 0x1E)

# Content area under the title (EMU) on the template's "Title and Content" layout.
X0, Y0, W0, H0 = 280556, 1330000, 11596254, 4950000
BODY_PT, SMALL_PT, CODE_PT = 12, 10, 9

# ------------------------------------------------------------ inline text --
INLINE = re.compile(r"(<b>.*?</b>|<strong>.*?</strong>|<i>.*?</i>|<em>.*?</em>|<code>.*?</code>|<a [^>]*>.*?</a>)", re.S)


def text_runs(par, h, size=BODY_PT, color=INK, bold=False, italic=False, font=None):
    """Add inline-HTML text (b / i / code / a, as produced by the deck renderer) to a paragraph."""
    for tok in INLINE.split(h or ""):
        if not tok:
            continue
        b, i, f, inner = bold, italic, font, tok
        m = re.match(r"<(b|strong|i|em|code|a)\b[^>]*>(.*)</\1>$", tok, re.S)
        if m:
            tag, inner = m.groups()
            b = b or tag in ("b", "strong")
            i = i or tag in ("i", "em")
            f = "Consolas" if tag == "code" else font
        txt = html.unescape(re.sub(r"<[^>]+>", "", inner))
        if not txt:
            continue
        r = par.add_run()
        r.text = txt
        r.font.size, r.font.bold, r.font.italic = Pt(size), b, i
        r.font.color.rgb = color
        if f:
            r.font.name = f


def md_inline(s):
    """Markdown inline (as in list items) -> the same inline HTML the deck renderer makes."""
    t = html.escape(s or "", quote=False)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<i>\1</i>", t)
    t = re.sub(r"`(.+?)`", r"<code>\1</code>", t)
    return t


def paragraphs(h):
    ps = re.findall(r"<p>(.*?)</p>", h or "", re.S)
    return ps or [h or ""]


def bullet(par, char="•"):
    pPr = par._p.get_or_add_pPr()
    pPr.set("marL", "228600")
    pPr.set("indent", "-228600")
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)


def numbered(par):
    pPr = par._p.get_or_add_pPr()
    pPr.set("marL", "342900")
    pPr.set("indent", "-342900")
    bu = etree.SubElement(pPr, qn("a:buAutoNum"))
    bu.set("type", "arabicPeriod")


def textbox(slide, rect, fill=None, line=None, inset=60000):
    x, y, w, h = rect
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE if fill or line else MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(inset))
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for el in list(bodyPr):
        bodyPr.remove(el)
    etree.SubElement(bodyPr, qn("a:normAutofit"))            # PowerPoint shrinks text that does not fit
    return shp, tf


def first_par(tf):
    return tf.paragraphs[0] if not tf.paragraphs[0].runs and len(tf.paragraphs) == 1 else tf.add_paragraph()


# ----------------------------------------------------------------- blocks --
def put_block(slide, b, rect, ctx):
    t = b["type"]
    if b.get("smartart") and (SMARTART / b["smartart"]).exists():
        return put_smartart(slide, b, rect, ctx)               # PowerPoint lays the diagram out inside rect
    if t in ("text", "callout", "pullquote"):
        fill = RED_SOFT if t == "callout" else None
        _, tf = textbox(slide, rect, fill=fill, inset=90000 if fill else 30000)
        for ph in paragraphs(b["html"]):
            p = first_par(tf)
            p.space_after = Pt(6)
            text_runs(p, ph, size=(BODY_PT + 3 if t == "pullquote" else BODY_PT), italic=t == "pullquote",
                      color=(RED_DEEP if t == "pullquote" else INK))
        return
    if t in ("list", "steps"):
        boxed = bool(b.get("heading"))
        _, tf = textbox(slide, rect, line=LINE if boxed else None, inset=90000 if boxed else 30000)
        if b.get("heading"):
            p = first_par(tf)
            text_runs(p, html.escape(b["heading"]), size=BODY_PT + 1, bold=True, color=RED_DEEP)
            p.space_after = Pt(4)
        if b.get("style") == "chips":
            p = first_par(tf)
            text_runs(p, "   ·   ".join(md_inline(i) for i in b["items"]), size=BODY_PT, color=RED_DEEP)
            return
        for item in b["items"]:
            p = first_par(tf)
            p.space_after = Pt(3)
            if t == "steps" or b.get("ordered"):
                numbered(p)
                item = re.sub(r"^\d+[.)]\s*", "", item)
            else:
                bullet(p)
            text_runs(p, md_inline(item))
        return
    if t == "flow":
        x, y, w, h = rect
        n = len(b["steps"])
        gap = 150000
        bw = (w - gap * (n - 1)) // max(n, 1)
        for k, s in enumerate(b["steps"]):
            shp = slide.shapes.add_shape(MSO_SHAPE.CHEVRON if k else MSO_SHAPE.PENTAGON,
                                         Emu(x + k * (bw + gap)), Emu(y), Emu(bw), Emu(max(h, 500000)))
            shp.fill.solid()
            shp.fill.fore_color.rgb = RED if k % 2 == 0 else RED_DEEP
            shp.line.fill.background()
            tf = shp.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            text_runs(p, html.escape(s), size=SMALL_PT, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
        return
    if t == "table":
        x, y, w, h = rect
        rows, cols = len(b["rows"]) + 1, max(len(b["headers"]), max((len(r) for r in b["rows"]), default=1))
        gt = slide.shapes.add_table(rows, cols, Emu(x), Emu(y), Emu(w), Emu(h)).table
        for r, row in enumerate([b["headers"]] + b["rows"]):
            for c in range(cols):
                cell = gt.cell(r, c)
                cell.margin_left = cell.margin_right = Emu(50000)
                cell.margin_top = cell.margin_bottom = Emu(25000)
                p = cell.text_frame.paragraphs[0]
                txt = row[c] if c < len(row) else ""
                text_runs(p, md_inline(txt) if r else html.escape(txt), size=SMALL_PT, bold=r == 0,
                          color=RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else INK)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RED if r == 0 else (RGBColor(0xFF, 0xFF, 0xFF) if r % 2 else RGBColor(0xFB, 0xFA, 0xF8))
        return
    if t == "code":
        _, tf = textbox(slide, rect, fill=CODE_BG, inset=90000)
        lines = b["code"].split("\n")
        for k, line in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = line or " "
            r.font.name, r.font.size = "Consolas", Pt(CODE_PT)
            r.font.color.rgb = RGBColor(0xD4, 0xD4, 0xD4)
        return
    if t == "image":
        src = ctx["image_path"](b["src"])
        if src is None:
            return
        x, y, w, h = rect
        cap_h = 260000 if b.get("caption") else 0
        from PIL import Image
        with Image.open(src) as im:
            iw, ih = im.size
        scale = min(w / iw, (h - cap_h) / ih)
        pw, ph = int(iw * scale), int(ih * scale)
        slide.shapes.add_picture(str(src), Emu(x + (w - pw) // 2), Emu(y), Emu(pw), Emu(ph))
        if b.get("caption"):
            _, tf = textbox(slide, (x, y + ph + 40000, w, cap_h), inset=0)
            p = first_par(tf)
            p.alignment = PP_ALIGN.CENTER
            text_runs(p, md_inline(b["caption"]), size=8, color=GREY)
        return
    if t == "card":
        x, y, w, h = rect
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
        shp.adjustments[0] = 0.04
        shp.fill.background()
        shp.line.color.rgb = LINE
        shp.line.width = Pt(1.25)
        inner_y = y + 70000
        if b.get("heading"):
            _, tf = textbox(slide, (x + 60000, inner_y, w - 120000, 330000), inset=30000)
            text_runs(first_par(tf), html.escape(b["heading"]), size=BODY_PT + 1, bold=True, color=RED_DEEP)
            inner_y += 330000
        rest = (x + 60000, inner_y, w - 120000, y + h - inner_y - 50000)
        stack(slide, b["blocks"], rest, ctx)
        return
    if t == "columns":
        x, y, w, h = rect
        n = len(b["columns"])
        gap = 170000
        cw = (w - gap * (n - 1)) // n
        for k, col in enumerate(b["columns"]):
            stack(slide, col, (x + k * (cw + gap), y, cw, h), ctx)
        return


def est(b):
    """Relative height of a block when no browser measurement is available."""
    t = b["type"]
    if t == "image":
        return 6.0
    if t == "columns":
        return max(sum(est(x) for x in c) for c in b["columns"])
    if t == "card":
        return 0.5 + sum(est(x) for x in b["blocks"])
    if t == "table":
        return 0.5 + 0.45 * len(b["rows"])
    if t == "code":
        return 0.3 + 0.22 * (b["code"].count("\n") + 1)
    if t in ("list", "steps"):
        return 0.4 + sum(0.35 + len(i) / 160 for i in b["items"])
    if t == "flow":
        return 1.2
    return 0.3 + len(re.sub(r"<[^>]+>", "", b.get("html", ""))) / 160


def stack(slide, blocks, rect, ctx):
    x, y, w, h = rect
    weights = [est(b) for b in blocks]
    total = sum(weights) or 1
    gap = 90000
    avail = h - gap * (len(blocks) - 1)
    for b, wt in zip(blocks, weights):
        bh = int(avail * wt / total)
        put_block(slide, b, (x, y, w, bh), ctx)
        y += bh + gap


# --------------------------------------------------------------- SmartArt --
DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
CT = "application/vnd.openxmlformats-officedocument.drawingml."
# How many text levels each layout shows, and whether the block heading is its centre node.
SMARTART_SHAPE = {"RadialCluster": (2, True), "radial6": (1, True), "HexagonCluster": (1, False),
                  "CircleProcess": (1, False), "VaryingWidthList": (1, False), "process2": (1, False),
                  "process5": (1, False), "vList5": (2, False), "chevron2": (2, False),
                  "VerticalCurvedList": (1, False)}


def smartart_tree(b):
    """Block items -> [(runs, children)], runs = [(text, bold)]. "**Lead:** rest" items become a node
    with a child (two-level layouts) or one node with a bold lead (one-level layouts)."""
    layout = b["smartart"]
    depth, rooted = SMARTART_SHAPE.get(layout, (1, False))
    items = b.get("items") or b.get("steps") or []
    nodes = []
    for it in items:
        it = re.sub(r"^\d+[.)]\s*", "", it)
        m = re.match(r"^\*\*(.+?)\*\*\s*:?\s*(.*)$", it)
        lead = rest = None
        if m and (m.group(1).endswith(":") or it[m.end(1) + 2:].lstrip().startswith(":")):
            lead, rest = m.group(1).rstrip(":").strip(), m.group(2).lstrip(":").strip()
        plain = lambda s: html.unescape(re.sub(r"[*`]", "", s))
        if lead and rest and depth >= 2:
            nodes.append(([(plain(lead), False)], [([(plain(rest), False)], [])]))
        elif lead and rest:
            nodes.append(([(plain(lead) + ": ", True), (plain(rest), False)], []))
        else:
            nodes.append(([(plain(it), False)], []))
    if rooted:
        nodes = [([(b.get("heading") or "", False)], nodes)]
    return nodes


def smartart_data(tree, prset):
    # Model ids derived from the content, so rebuilding an unchanged deck gives an identical file.
    seed, counter = json.dumps([tree, prset], sort_keys=True, default=str), iter(range(1 << 30))
    gid = lambda: "{" + str(uuid.uuid5(uuid.NAMESPACE_URL, f"{seed}#{next(counter)}")).upper() + "}"
    E = lambda tag, **kw: etree.Element(f"{{{DGM}}}{tag}", **kw)

    def textel(parent, runs):
        t = etree.SubElement(parent, f"{{{DGM}}}t")
        etree.SubElement(t, f"{{{A}}}bodyPr")
        etree.SubElement(t, f"{{{A}}}lstStyle")
        p = etree.SubElement(t, f"{{{A}}}p")
        for txt, bold in runs:
            r = etree.SubElement(p, f"{{{A}}}r")
            rp = etree.SubElement(r, f"{{{A}}}rPr", lang="en-US")
            if bold:
                rp.set("b", "1")
            etree.SubElement(r, f"{{{A}}}t").text = txt
        if not runs:
            etree.SubElement(p, f"{{{A}}}endParaRPr", lang="en-US")

    root = etree.Element(f"{{{DGM}}}dataModel", nsmap={"dgm": DGM, "a": A})
    pts = etree.SubElement(root, f"{{{DGM}}}ptLst")
    cxns = etree.SubElement(root, f"{{{DGM}}}cxnLst")
    doc = gid()
    d = etree.SubElement(pts, f"{{{DGM}}}pt", modelId=doc, type="doc")
    etree.SubElement(d, f"{{{DGM}}}prSet", **{k: v for k, v in prset.items() if k != "phldr"})
    etree.SubElement(d, f"{{{DGM}}}spPr")
    textel(d, [])

    def add(parent, nodes):
        for order, (runs, kids) in enumerate(nodes):
            nid, par, sib, cx = gid(), gid(), gid(), gid()
            n = etree.SubElement(pts, f"{{{DGM}}}pt", modelId=nid)
            etree.SubElement(n, f"{{{DGM}}}prSet")
            etree.SubElement(n, f"{{{DGM}}}spPr")
            textel(n, runs)
            for kind, tid in (("parTrans", par), ("sibTrans", sib)):
                tp = etree.SubElement(pts, f"{{{DGM}}}pt", modelId=tid, type=kind, cxnId=cx)
                etree.SubElement(tp, f"{{{DGM}}}prSet")
                etree.SubElement(tp, f"{{{DGM}}}spPr")
                textel(tp, [])
            etree.SubElement(cxns, f"{{{DGM}}}cxn", modelId=cx, srcId=parent, destId=nid, srcOrd=str(order),
                             destOrd="0", parTransId=par, sibTransId=sib)
            add(nid, kids)

    add(doc, tree)
    etree.SubElement(root, f"{{{DGM}}}bg")
    etree.SubElement(root, f"{{{DGM}}}whole")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def put_smartart(slide, b, rect, ctx):
    folder = SMARTART / b["smartart"]
    prset = json.loads((folder / "prset.json").read_text(encoding="utf-8"))
    ctx["n_diagrams"] += 1
    n = ctx["n_diagrams"]
    pkg = slide.part.package
    rels = {}
    for key, name, kind, ctype, blob in (
            ("dm", "data", "diagramData", "diagramData+xml", smartart_data(smartart_tree(b), prset)),
            ("lo", "layout", "diagramLayout", "diagramLayout+xml", (folder / "layout.xml").read_bytes()),
            ("qs", "quickStyle", "diagramQuickStyle", "diagramStyle+xml", (folder / "quickStyle.xml").read_bytes()),
            ("cs", "colors", "diagramColors", "diagramColors+xml", (folder / "colors.xml").read_bytes())):
        part = Part(PackURI(f"/ppt/diagrams/{name}{n}.xml"), CT + ctype, pkg, blob)
        rels[key] = slide.part.relate_to(part, RT + kind)
    x, y, w, h = rect
    sid = max([int(s.shape_id) for s in slide.shapes] + [1]) + 1
    gf = etree.fromstring(
        f'<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="{A}" xmlns:r="{R}">'
        f'<p:nvGraphicFramePr><p:cNvPr id="{sid}" name="Diagram {sid}"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
        f'<p:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></p:xfrm>'
        f'<a:graphic><a:graphicData uri="{DGM}"><dgm:relIds xmlns:dgm="{DGM}" r:dm="{rels["dm"]}" r:lo="{rels["lo"]}" '
        f'r:qs="{rels["qs"]}" r:cs="{rels["cs"]}"/></a:graphicData></a:graphic></p:graphicFrame>')
    slide.shapes._spTree.append(gf)


# ------------------------------------------------------------------ deck --
def notes_text(h):
    h = re.sub(r"</p>\s*<p>", "\n\n", h or "")
    h = re.sub(r"<li>", "\n• ", h)
    return html.unescape(re.sub(r"<[^>]+>", "", h)).strip()


def layout_named(prs, name):
    return next(l for l in prs.slide_layouts if l.name == name)


def build(pages, deck, layouts, image_path, out_path):
    """pages: the browser deck's pages (dicts with title/unit/blocks/notes/cover);
    layouts: per page, measured block rects in browser px ({"body": [w, h], "blocks": [[x, y, w, h], ...]})
    or None to stack blocks by estimate."""
    prs = Presentation(str(TEMPLATES / "lecture.pptx"))
    ctx = {"image_path": image_path, "n_diagrams": 0}
    for page, lay in zip(pages, layouts):
        if page.get("cover"):
            s = prs.slides.add_slide(layout_named(prs, "1_Title Slide"))
            s.shapes.title.text = re.sub(r"<[^>]+>", "", page["title"])
            _, tf = textbox(s, (655315, 1553437, 10515600, 420000), inset=0)
            text_runs(first_par(tf), html.escape(str(page.get("unit") or deck.get("kicker", "")).upper()),
                      size=14, bold=True, color=RED)
            # The title layout has artwork below the title, so the cover keeps to its subtitle line
            # (the browser deck's cover also shows the overview chips and callout).
            sub = next((b for b in page["blocks"] if b["type"] == "text"), None)
            if sub:
                _, tf = textbox(s, (655315, 3450000, 6200000, 700000), inset=0)
                text_runs(first_par(tf), paragraphs(sub["html"])[0], size=16, color=GREY)
        else:
            s = prs.slides.add_slide(layout_named(prs, "Title and Content"))
            s.shapes.title.text = re.sub(r"<[^>]+>", "", page["title"])
            for ph in list(s.placeholders):
                if ph.placeholder_format.idx == 1:                 # we place our own shapes in the body area
                    ph._element.getparent().remove(ph._element)
            if lay and len(lay["blocks"]) == len(page["blocks"]):
                bw, bh = lay["body"]
                bottom = max([r[1] + r[3] for r in lay["blocks"]] + [bh])
                sc = min(W0 / bw, H0 / bottom)
                for b, (x, y, w, h) in zip(page["blocks"], lay["blocks"]):
                    put_block(s, b, (X0 + int(x * sc), Y0 + int(y * sc), int(w * sc), int(h * sc)), ctx)
            else:
                stack(s, page["blocks"], (X0, Y0, W0, H0), ctx)
        if page.get("notes"):
            s.notes_slide.notes_text_frame.text = notes_text(page["notes"])
    prs.save(out_path if hasattr(out_path, "write") else str(out_path))
    return len(prs.slides)
